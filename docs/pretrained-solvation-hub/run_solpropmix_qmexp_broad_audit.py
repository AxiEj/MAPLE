from __future__ import annotations

import argparse
import contextlib
import csv
import hashlib
import importlib
import io
import json
import math
import subprocess
import sys
import tempfile
import time
from pathlib import Path
from string import ascii_uppercase
from typing import Any

import numpy as np
import pandas as pd
import torch
from openpyxl import load_workbook
from rdkit import Chem, RDLogger

REPO = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO))

from maple.function.benchmarking.pretrained_hub import (
    derive_functional_group_assignment,
    load_functional_group_taxonomy,
)
from maple.function.solvfe.solpropmix_property import (
    SOLPROPMIX_CHECKPOINT_FAMILY,
    SOLPROPMIX_CHECKPOINT_ORIGIN,
    SOLPROPMIX_CHECKPOINT_SHA256,
    SOLPROPMIX_CHECKPOINTS_BYTE_IDENTICAL_ACROSS_V1_0_V1_1,
    SOLPROPMIX_DATA_RELEASE_RECORD,
    SOLPROPMIX_DATA_RELEASE_VERSION,
    SOLPROPMIX_MODELWEIGHTS_ZIP_MD5,
    SOLPROPMIX_MODELWEIGHTS_ZIP_SHA256,
    SOLPROPMIX_MODELWEIGHTS_ZIP_SIZE_BYTES,
    SOLPROPMIX_PREVIOUS_DATA_RELEASE_RECORD,
    SOLPROPMIX_SOURCE_REVISION,
    SOLPROPMIX_SOURCE_TREE,
    SOLPROPMIX_SOURCE_URL,
    SOLPROPMIX_STATIC_CODE_ZIP_MD5,
    SOLPROPMIX_STATIC_CODE_ZIP_NAME,
    SOLPROPMIX_STATIC_CODE_ZIP_SHA256,
    SOLPROPMIX_STATIC_CODE_ZIP_SIZE_BYTES,
    SOLPROPMIX_SUPPLEMENTAL_SOURCE_ORIGIN,
    SolPropMixPropertyAdapter,
    _canonical_neutral_structure,
    _normalize_origin,
    _normalize_solvents,
)

RDLogger.DisableLog("rdApp.*")  # type: ignore[attr-defined]

WORKBOOK: Path
SOURCE_ROOT: Path
STATIC_ROOT: Path
WEIGHTS_ROOT: Path
PYTHON_EXE: Path
TAXONOMY = REPO / "docs/pretrained-solvation-hub/functional-group-taxonomy-v1.json"
TAXONOMY_SHA256 = "46e073d0e253d7534b7f516f1cc2096ba0abfea84ddc170c95fc86634c917b36"
WORKBOOK_SHA256 = "793f325e6edddca3334b368f68edfc368027e472ad779e7e5d52ff5d1e1a001d"
GPU_WORSEN_TOL = 5e-12
MAIN_SAMPLE_N = 1000
TERNARY_SAMPLE_N = 100
OUT_DIR: Path
DataPoint: Any
DatapointList: Any
DataTensor: Any
MolencoderDatabase: Any
GsolvHsolv: Any
load_checkpoint: Any
load_scaler: Any
STATIC_RUNTIME: tempfile.TemporaryDirectory[str] | None = None
DATA_RUNTIME: tempfile.TemporaryDirectory[str] | None = None
FROZEN_RUNTIME: tempfile.TemporaryDirectory[str] | None = None

MAIN_SHEETS = ("From IDAC - no water", "From VLE - Bin Solv - no water")
TERNARY_SHEET = "From VLE - Ter Solv - no water"
OFFICIAL_STATIC_PYTHON_SHA256 = {
    "setup.py": "4df8465b40cf6e51f37a281eeab76f15ffcdb096d646d91b0788c0d8c418e6b1",
    "solvation_predictor/__init__.py": "e3b0c44298fc1c149afbf4c8996fb92427ae41e4649b934ca495991b7852b855",
    "solvation_predictor/data/Scaler.py": "5e4a65084db838238f14386af66672184d93764764b694b138fa97cf13dffa73",
    "solvation_predictor/data/Splitter.py": "f5244d923b5afec61e22d6d9ded1bef081545415b9b0efefe7282d8ab5e7d1ac",
    "solvation_predictor/data/__init__.py": "2639b97ff1fe63a3c9e88402aa9d1369305be046762483db834325da22948a78",
    "solvation_predictor/data/data.py": "fd7bc4454b981686b3ef7fbb86d5069f9cb5eb7b4bbc454e3876ac61775aa83e",
    "solvation_predictor/features/MolEncoder.py": "8870322c62ee162c3c378eb2800587e3710188fd111ab1fffc12525f32aa91e2",
    "solvation_predictor/features/__init__.py": "e3b0c44298fc1c149afbf4c8996fb92427ae41e4649b934ca495991b7852b855",
    "solvation_predictor/features/calculated_features.py": "55300946dd269733de164ef75c1265f5b31243935d85104ea08377198f1cf29f",
    "solvation_predictor/features/feature_vectors.py": "3b4d7e30fd48f9f2d213f936c10ca98a26290875f80178fa0f797187a39fded8",
    "solvation_predictor/inp.py": "e46b809a992a382794858d7cdc6406cff017ac4b37fd5fda92b27055d5eca167",
    "solvation_predictor/models/Concat.py": "fd914b664c29d748790d9fc1b0951f122c6557242021722679fb6de23f7d2898",
    "solvation_predictor/models/FFN.py": "fafe31579a68d7e0a6881b0600279aba5b69a36cafd471f526e732a97e30efbc",
    "solvation_predictor/models/MPN.py": "6f445b5860be280048f1c61567e7a0183f742419a83b767c616a9fdddf42a7f0",
    "solvation_predictor/models/MPNNconv.py": "5c2d59464ebdc2d1285b7d8e2b376031e149d9a83ca7ff181be8862d8f80cade",
    "solvation_predictor/models/Model.py": "959160d54628acdce76dabe573e4f857b4796d87f5c2c75e23232bfb3c455a9a",
    "solvation_predictor/models/__init__.py": "e3b0c44298fc1c149afbf4c8996fb92427ae41e4649b934ca495991b7852b855",
    "solvation_predictor/models/solvGNN.py": "ad7c31489069c285ce4b240e26d3e12869b9d7cd562fbb7a196f17b0d58acd38",
    "solvation_predictor/predict.py": "58e982f203879a2278e37e616fa708f9a7e36f7d112e2940133ef2150454a49c",
    "solvation_predictor/solubility/SolubilityCalculations.py": "58fef5e69288da458f9a5a8c1b59837663dbcc1f47d8b26f9350cb684a0903f3",
    "solvation_predictor/solubility/SolubilityData.py": "3e4e941ceb19d951d3e736b9e31eeaa102c47088b50bff316745f622f1537294",
    "solvation_predictor/solubility/SolubilityModels.py": "3fe8ebb9f2303c1cb98b8a447a34d207f0c0194ab50d954aa13f7bb9cac12718",
    "solvation_predictor/solubility/SolubilityPredictions.py": "1d199488865ebd3c27624af87ea05455194a7756fef3ff1ddeebe4e1bbe9be16",
    "solvation_predictor/solubility/__init__.py": "e3b0c44298fc1c149afbf4c8996fb92427ae41e4649b934ca495991b7852b855",
    "solvation_predictor/solubility/calculate_solubility.py": "d83f150892c7f5c5f8fc577242e52d4aa128da3a8f249f03411b76179fd194e6",
    "solvation_predictor/solubility_predict.py": "f5c119d8be3c7b363362eeed33083464978be7cebabfb058fd1a50717e846eeb",
    "solvation_predictor/train.py": "34143553b309a083ee72fda3b9b203a2ea6e5e8de08878a2367f857d2f72c33e",
    "solvation_predictor/train/__init__.py": "e3b0c44298fc1c149afbf4c8996fb92427ae41e4649b934ca495991b7852b855",
    "solvation_predictor/train/evaluate.py": "e2e023e8f44a9c958b9986029ec0626f091ae97bea16aa25cdd8df582e5a2eae",
    "solvation_predictor/train/train.py": "556fd9b8ed5dcd531d2e58b845ecc95fb43357faf31c28c7863a7a396cd8cab6",
    "solvation_predictor/train/utils.py": "b00eb3d1ffd4ff9d5472aa17536a9f06714ab5166dc6cdfe7820b69163810957",
}
CANONICAL_FROZEN_INPUT_SHA256 = {
    "adapter-crosscheck.json": "344657e8da67945a5c2fde3e6b7b877401a64631317d1e69809ea813327d1386",
    "experimental-references.json": "96ae81e15db0ed5cdbafdc85f6d7d2ba53cebde8ab2cf13d34f1f82bbafe6906",
    "panel-inputs.json": "6e593c52eadcc6f1d3793161e80a30e31fe5bb83e21dd0fa1379db2bbb43bd24",
    "rows.csv": "626f4c7f1372b5036360654a7efb9b158eed00fda2809efc823f5d1947d82084",
    "ternary-rows.csv": "87535555c65fead1f32e15e5a0f09e6eabacc5280754ef0bd73166f6c89dad8a",
}


def _git_output(root: Path, *args: str) -> str:
    result = subprocess.run(
        ["git", "-C", str(root), *args],
        check=False,
        capture_output=True,
        text=True,
    )
    if result.returncode != 0:
        detail = result.stderr.strip() or result.stdout.strip()
        raise RuntimeError(f"unable to authenticate source checkout: {detail}")
    return result.stdout.strip()


def authenticate_official_assets(
    workbook: Path,
    source_root: Path,
    static_root: Path,
    weights_root: Path,
    python_executable: Path,
) -> None:
    """Fail closed on every executable/input identity before upstream import/load."""
    for root, label in (
        (source_root, "source root"),
        (static_root, "static root"),
        (weights_root, "weights root"),
    ):
        if not root.is_dir() or root.is_symlink():
            raise RuntimeError(f"{label} must be a real directory")
    if not workbook.is_file() or workbook.is_symlink():
        raise RuntimeError("workbook must be a real file")
    try:
        resolved_python = python_executable.resolve(strict=True)
    except OSError as exc:
        raise RuntimeError("python executable does not resolve to a file") from exc
    if not resolved_python.is_file():
        raise RuntimeError("python executable does not resolve to a regular file")
    if sha256_file(workbook) != WORKBOOK_SHA256:
        raise RuntimeError("workbook identity mismatch")
    if _git_output(source_root, "rev-parse", "--is-inside-work-tree") != "true":
        raise RuntimeError("source root is not a git work tree")
    if _git_output(source_root, "rev-parse", "HEAD") != SOLPROPMIX_SOURCE_REVISION:
        raise RuntimeError("source revision mismatch")
    if _git_output(source_root, "rev-parse", "HEAD^{tree}") != SOLPROPMIX_SOURCE_TREE:
        raise RuntimeError("source tree mismatch")
    origin = _git_output(source_root, "config", "--get", "remote.origin.url")
    if _normalize_origin(origin) != _normalize_origin(SOLPROPMIX_SOURCE_URL):
        raise RuntimeError("source origin mismatch")
    if _git_output(source_root, "status", "--porcelain", "--untracked-files=all"):
        raise RuntimeError("source checkout is not clean")

    observed_python = {
        path.relative_to(static_root).as_posix()
        for path in static_root.rglob("*.py")
        if path.is_file()
    }
    if observed_python != set(OFFICIAL_STATIC_PYTHON_SHA256):
        raise RuntimeError("static Python module set mismatch")
    for relative, expected in OFFICIAL_STATIC_PYTHON_SHA256.items():
        path = static_root / relative
        if path.is_symlink() or sha256_file(path) != expected:
            raise RuntimeError(f"static module identity mismatch: {relative}")

    expected_names = [f"model{index}.pt" for index in range(10)]
    observed_names = sorted(path.name for path in weights_root.iterdir())
    if observed_names != expected_names:
        raise RuntimeError(
            "weights root must contain exactly model0.pt through model9.pt"
        )
    for name, expected in zip(
        expected_names, SOLPROPMIX_CHECKPOINT_SHA256, strict=True
    ):
        path = weights_root / name
        if path.is_symlink() or sha256_file(path) != expected:
            raise RuntimeError(f"checkpoint identity mismatch: {name}")


def snapshot_official_data(workbook: Path, weights_root: Path) -> tuple[Path, Path]:
    """Copy authenticated data to private paths and rehash every destination."""
    global DATA_RUNTIME
    if DATA_RUNTIME is not None:
        raise RuntimeError("official data snapshot is already configured")
    runtime = tempfile.TemporaryDirectory(prefix="solpropmix-data-")
    try:
        private_root = Path(runtime.name)
        private_workbook = private_root / "solprop-mix_v1.1.xlsx"
        private_workbook.write_bytes(workbook.read_bytes())
        if sha256_file(private_workbook) != WORKBOOK_SHA256:
            raise RuntimeError("private workbook snapshot identity mismatch")
        private_weights = private_root / "weights"
        private_weights.mkdir()
        for index, expected in enumerate(SOLPROPMIX_CHECKPOINT_SHA256):
            name = f"model{index}.pt"
            destination = private_weights / name
            destination.write_bytes((weights_root / name).read_bytes())
            if sha256_file(destination) != expected:
                raise RuntimeError(
                    f"private checkpoint snapshot identity mismatch: {name}"
                )
    except Exception:
        runtime.cleanup()
        raise
    DATA_RUNTIME = runtime
    return private_workbook, private_weights


def snapshot_canonical_frozen_inputs(frozen_dir: Path) -> Path:
    """Authenticate scored bytes once, then use only a private snapshot."""
    global FROZEN_RUNTIME
    if FROZEN_RUNTIME is not None:
        raise RuntimeError("canonical frozen-input snapshot is already configured")
    runtime = tempfile.TemporaryDirectory(prefix="solpropmix-frozen-")
    try:
        private_root = Path(runtime.name)
        for name, expected in CANONICAL_FROZEN_INPUT_SHA256.items():
            source = frozen_dir / name
            if not source.is_file() or source.is_symlink():
                raise RuntimeError(f"canonical frozen input is not a real file: {name}")
            data = source.read_bytes()
            if hashlib.sha256(data).hexdigest() != expected:
                raise RuntimeError(f"canonical frozen input identity mismatch: {name}")
            destination = private_root / name
            destination.write_bytes(data)
            if sha256_file(destination) != expected:
                raise RuntimeError(
                    f"private canonical frozen input identity mismatch: {name}"
                )
    except Exception:
        runtime.cleanup()
        raise
    FROZEN_RUNTIME = runtime
    return private_root


def configure_runtime(static_root: Path) -> None:
    """Import an authenticated source-only copy, never the supplied tree."""
    global DataPoint, DatapointList, DataTensor, MolencoderDatabase
    global GsolvHsolv, load_checkpoint, load_scaler, STATIC_RUNTIME
    if STATIC_RUNTIME is not None:
        raise RuntimeError("supplemental runtime is already configured")
    STATIC_RUNTIME = tempfile.TemporaryDirectory(prefix="solpropmix-static-")
    clean_root = Path(STATIC_RUNTIME.name)
    for relative, expected in OFFICIAL_STATIC_PYTHON_SHA256.items():
        source = static_root / relative
        destination = clean_root / relative
        destination.parent.mkdir(parents=True, exist_ok=True)
        destination.write_bytes(source.read_bytes())
        if destination.is_symlink() or sha256_file(destination) != expected:
            raise RuntimeError(f"clean static copy identity mismatch: {relative}")
    sys.path.insert(0, str(clean_root))
    data_module = importlib.import_module("solvation_predictor.data.data")
    inp_module = importlib.import_module("solvation_predictor.inp")
    train_module = importlib.import_module("solvation_predictor.train.train")
    DataPoint = data_module.DataPoint
    DatapointList = data_module.DatapointList
    DataTensor = data_module.DataTensor
    MolencoderDatabase = data_module.MolencoderDatabase
    GsolvHsolv = inp_module.GsolvHsolv
    load_checkpoint = train_module.load_checkpoint
    load_scaler = train_module.load_scaler


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def supported_structure(value: Any, role: str) -> bool:
    if not isinstance(value, str) or not value.startswith("InChI="):
        return False
    try:
        _canonical_neutral_structure(value, role)
        return True
    except Exception:  # noqa: BLE001 - upstream canonicalizer has a broad API
        return False


def metrics(errors: list[float]) -> dict[str, float | int]:
    return {
        "n": len(errors),
        "mae_kcal_mol": math.fsum(abs(x) for x in errors) / len(errors),
        "rmse_kcal_mol": math.sqrt(math.fsum(x * x for x in errors) / len(errors)),
        "maxae_kcal_mol": max(abs(x) for x in errors),
    }


def header_map(sheet: str) -> dict[str, str]:
    wb = load_workbook(WORKBOOK, read_only=True, data_only=True)
    ws = wb[sheet]
    header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    out: dict[str, str] = {}
    for idx, name in enumerate(header, start=1):
        col = ""
        n = idx
        while n:
            n, rem = divmod(n - 1, 26)
            col = ascii_uppercase[rem] + col
        out[str(name)] = col
    wb.close()
    return out


def strict_precision_flags() -> None:
    torch.use_deterministic_algorithms(True)
    torch.set_float32_matmul_precision("highest")
    torch.backends.cudnn.allow_tf32 = False
    torch.backends.cudnn.benchmark = False
    torch.backends.cudnn.deterministic = True
    torch.backends.cuda.matmul.allow_tf32 = False
    if hasattr(torch.backends.cuda.matmul, "allow_fp16_reduced_precision_reduction"):
        torch.backends.cuda.matmul.allow_fp16_reduced_precision_reduction = False
    if hasattr(torch.backends.cuda.matmul, "allow_bf16_reduced_precision_reduction"):
        torch.backends.cuda.matmul.allow_bf16_reduced_precision_reduction = False


def build_models(device_name: str):
    strict_precision_flags()
    torch.set_default_dtype(torch.float64)
    inp = GsolvHsolv().parse_args([])
    inp.cuda = device_name.startswith("cuda")
    inp.device = torch.device(device_name)
    bundles = []
    t0 = time.perf_counter()
    for i in range(10):
        p = WEIGHTS_ROOT / f"model{i}.pt"
        scaler = load_scaler(str(p), from_package=False)
        with contextlib.redirect_stdout(io.StringIO()):
            model = load_checkpoint(str(p), inp, from_package=False)
        model = model.to(device=inp.device, dtype=torch.float64)
        model.eval()
        bundles.append((model, scaler))
    receipt = {
        "execution_device": device_name,
        "execution_dtype": "float64",
        "torch_version": torch.__version__,
        "cuda_version": torch.version.cuda,
        "rdkit_version": Chem.rdBase.rdkitVersion,
        "numpy_version": np.__version__,
        "deterministic_algorithms": bool(torch.are_deterministic_algorithms_enabled()),
        "float32_matmul_precision": torch.get_float32_matmul_precision(),
        "cuda_matmul_allow_tf32": bool(torch.backends.cuda.matmul.allow_tf32),
        "cudnn_allow_tf32": bool(torch.backends.cudnn.allow_tf32),
        "amp_autocast_enabled": False,
        "fp16_reduced_precision_reduction": bool(
            getattr(
                torch.backends.cuda.matmul,
                "allow_fp16_reduced_precision_reduction",
                False,
            )
        ),
        "bf16_reduced_precision_reduction": bool(
            getattr(
                torch.backends.cuda.matmul,
                "allow_bf16_reduced_precision_reduction",
                False,
            )
        ),
        "build_seconds": time.perf_counter() - t0,
    }
    return inp, bundles, receipt


def build_tensors(data: Any, inp: Any) -> list[Any]:
    grouped = [[] for _ in range(inp.max_num_mols)]
    for mol_index in range(inp.max_num_mols):
        for d in data.get_data():
            encoders = d.get_mol_encoder()
            if len(encoders) < inp.max_num_mols:
                encoders = encoders + [encoders[0]] * (inp.max_num_mols - len(encoders))
            grouped[mol_index].append(encoders[mol_index])
    out = []
    for group in grouped:
        tensor = DataTensor(group, inp, property=inp.property)
        tensor.f_atoms = tensor.f_atoms.to(device=inp.device, dtype=torch.float64)
        tensor.f_bonds = tensor.f_bonds.to(device=inp.device, dtype=torch.float64)
        tensor.f_mols = tensor.f_mols.to(device=inp.device, dtype=torch.float64)
        tensor.a2b = tensor.a2b.to(device=inp.device)
        tensor.b2a = tensor.b2a.to(device=inp.device)
        tensor.b2revb = tensor.b2revb.to(device=inp.device)
        out.append(tensor)
    return out


def prepare_request(solute_inchi: str, solvent_pairs: list[tuple[str, float]]):
    canonical_solute = _canonical_neutral_structure(solute_inchi, "solute")
    canonical_solvents = _normalize_solvents(solvent_pairs)
    request_solvents = [
        {"smiles": c.canonical_smiles, "fraction": c.mole_fraction}
        for c in canonical_solvents
    ]
    return canonical_solute, canonical_solvents, request_solvents


def predict_direct(
    inp_template: Any,
    bundles,
    *,
    solute_inchi: str,
    solvent_pairs: list[tuple[str, float]],
    temperature_k: float,
):
    canonical_solute, canonical_solvents, request_solvents = prepare_request(
        solute_inchi, solvent_pairs
    )
    inp = GsolvHsolv().parse_args([])
    inp.cuda = inp_template.cuda
    inp.device = inp_template.device
    inp.max_num_mols = len(request_solvents) + 1
    point = DataPoint(
        [canonical_solute] + [x["smiles"] for x in request_solvents],
        [0.0, 0.0],
        [],
        [x["fraction"] for x in request_solvents],
        inp,
        MolencoderDatabase(),
    )
    data = DatapointList([point])
    tensors = build_tensors(data, inp)
    preds = []
    for model_index, (model, scaler) in enumerate(bundles):
        model.inp = inp
        model.max_num_mols = inp.max_num_mols
        with torch.no_grad():
            out = model(data, tensors).reshape(1, -1).detach().cpu().numpy()
        val = scaler.inverse_transform(out)[0]
        g298 = float(val[0])
        h298 = float(val[1])
        gt = float(
            temperature_k
            * (g298 / 298.15 - h298 * (1.0 / 298.15 - 1.0 / temperature_k))
        )
        preds.append(
            {
                "model_index": model_index,
                "g298_kcal_mol": g298,
                "h298_kcal_mol": h298,
                "g_temperature_kcal_mol": gt,
            }
        )
    vals = [x["g_temperature_kcal_mol"] for x in preds]
    return {
        "canonical_solute_smiles": canonical_solute,
        "canonical_solvents": [
            {"canonical_smiles": c.canonical_smiles, "mole_fraction": c.mole_fraction}
            for c in canonical_solvents
        ],
        "prediction_kcal_mol": math.fsum(vals) / len(vals),
        "ensemble_population_std_kcal_mol": float(np.std(vals, ddof=0)),
        "model_predictions": preds,
    }


def make_row_record(
    sheet: str, row: Any, hmap: dict[str, str], selection_position: int
):
    excel_row = int(row["sheet_zero_based_row_index"]) + 2
    if sheet == "From IDAC - no water":
        solvent_pairs = [(str(row["inchi_solvent1"]), 1.0)]
        solvent_names = [str(row["DDB_name_solvent1"])]
    elif sheet == "From VLE - Bin Solv - no water":
        x = float(row["frac_solvent1"])
        solvent_pairs = [
            (str(row["inchi_solvent1"]), x),
            (str(row["inchi_solvent2"]), 1.0 - x),
        ]
        solvent_names = [str(row["DDB_name_solvent1"]), str(row["DDB_name_solvent2"])]
    else:
        x1 = float(row["frac_solvent1"])
        x2 = float(row["frac_solvent2"])
        solvent_pairs = [
            (str(row["inchi_solvent1"]), x1),
            (str(row["inchi_solvent2"]), x2),
            (str(row["inchi_solvent3"]), 1.0 - x1 - x2),
        ]
        solvent_names = [
            str(row["DDB_name_solvent1"]),
            str(row["DDB_name_solvent2"]),
            str(row["DDB_name_solvent3"]),
        ]
    return {
        "row_id": f"{sheet}!{excel_row}",
        "sheet": sheet,
        "excel_row_number": excel_row,
        "sheet_zero_based_row_index": int(row["sheet_zero_based_row_index"]),
        "selection_position": int(selection_position),
        "primary_functional_group": (
            str(row["primary_functional_group"])
            if "primary_functional_group" in row.index
            else None
        ),
        "all_matched_groups": (
            list(row["all_matched_groups"])
            if "all_matched_groups" in row.index
            else None
        ),
        "solute_name": str(row["DDB_name_solute"]),
        "solute_inchi": str(row["inchi_solute"]),
        "temperature_kelvin": float(row["T (K)"]),
        "experimental_gsolv_kcal_mol": (
            float(row["Gsolv (kcal/mol)"])
            if "Gsolv (kcal/mol)" in row.index and not pd.isna(row["Gsolv (kcal/mol)"])
            else None
        ),
        "published_workbook_qmexp_h298_kcal_mol": (
            float(row["SolProp-mix_QM_Exp_Hsolv298"])
            if "SolProp-mix_QM_Exp_Hsolv298" in row.index
            and not pd.isna(row["SolProp-mix_QM_Exp_Hsolv298"])
            else None
        ),
        "published_workbook_qmexp_g298_kcal_mol": (
            float(row["SolProp-mix_QM_Exp_Gsolv298"])
            if "SolProp-mix_QM_Exp_Gsolv298" in row.index
            and not pd.isna(row["SolProp-mix_QM_Exp_Gsolv298"])
            else None
        ),
        "published_workbook_qmexp_gt_kcal_mol": (
            float(row["SolProp-mix_QM_Exp_GsolvT"])
            if "SolProp-mix_QM_Exp_GsolvT" in row.index
            and not pd.isna(row["SolProp-mix_QM_Exp_GsolvT"])
            else None
        ),
        "solvent_pairs": solvent_pairs,
        "solvent_names": solvent_names,
        "input_cells": {
            "solute_inchi": f"{hmap['inchi_solute']}{excel_row}",
            "solute_name": f"{hmap['DDB_name_solute']}{excel_row}",
            "solvent1_inchi": f"{hmap['inchi_solvent1']}{excel_row}",
            "solvent1_name": f"{hmap['DDB_name_solvent1']}{excel_row}",
            "solvent2_inchi": (
                f"{hmap['inchi_solvent2']}{excel_row}"
                if "inchi_solvent2" in hmap and len(solvent_pairs) > 1
                else None
            ),
            "solvent2_name": (
                f"{hmap['DDB_name_solvent2']}{excel_row}"
                if "DDB_name_solvent2" in hmap and len(solvent_pairs) > 1
                else None
            ),
            "solvent3_inchi": (
                f"{hmap['inchi_solvent3']}{excel_row}"
                if "inchi_solvent3" in hmap and len(solvent_pairs) > 2
                else None
            ),
            "solvent3_name": (
                f"{hmap['DDB_name_solvent3']}{excel_row}"
                if "DDB_name_solvent3" in hmap and len(solvent_pairs) > 2
                else None
            ),
            "frac_solvent1": (
                f"{hmap['frac_solvent1']}{excel_row}"
                if "frac_solvent1" in hmap and len(solvent_pairs) > 1
                else None
            ),
            "frac_solvent2": (
                f"{hmap['frac_solvent2']}{excel_row}"
                if "frac_solvent2" in hmap and len(solvent_pairs) > 2
                else None
            ),
            "temperature_kelvin": f"{hmap['T (K)']}{excel_row}",
            "experimental_gsolv": (
                f"{hmap['Gsolv (kcal/mol)']}{excel_row}"
                if "Gsolv (kcal/mol)" in hmap
                else None
            ),
        },
    }


LABEL_COLUMNS = (
    "Gsolv (kcal/mol)",
    "SolProp-mix_QM_Exp_Hsolv298",
    "SolProp-mix_QM_Exp_Gsolv298",
    "SolProp-mix_QM_Exp_GsolvT",
)


def attach_labels_after_selection(rows: list[dict[str, Any]]) -> None:
    """Read labels only after the deterministic input-only row IDs are frozen."""
    by_sheet: dict[str, list[dict[str, Any]]] = {}
    for row in rows:
        by_sheet.setdefault(row["sheet"], []).append(row)
    for sheet, selected in by_sheet.items():
        labels = pd.read_excel(
            WORKBOOK,
            sheet_name=sheet,
            engine="openpyxl",
            usecols=list(LABEL_COLUMNS),
        )
        for row in selected:
            label_row = labels.iloc[row["sheet_zero_based_row_index"]]
            row["experimental_gsolv_kcal_mol"] = float(label_row[LABEL_COLUMNS[0]])
            row["published_workbook_qmexp_h298_kcal_mol"] = float(
                label_row[LABEL_COLUMNS[1]]
            )
            row["published_workbook_qmexp_g298_kcal_mol"] = float(
                label_row[LABEL_COLUMNS[2]]
            )
            row["published_workbook_qmexp_gt_kcal_mol"] = float(
                label_row[LABEL_COLUMNS[3]]
            )


def select_main_rows():
    taxonomy = load_functional_group_taxonomy(TAXONOMY, expected_sha256=TAXONOMY_SHA256)
    hm = {sheet: header_map(sheet) for sheet in MAIN_SHEETS}
    frames = []
    for sheet in MAIN_SHEETS:
        usecols = [
            "inchi_solute",
            "DDB_name_solute",
            "inchi_solvent1",
            "DDB_name_solvent1",
            "T (K)",
        ]
        if "Bin" in sheet:
            usecols += ["inchi_solvent2", "DDB_name_solvent2", "frac_solvent1"]
        df = (
            pd.read_excel(
                WORKBOOK, sheet_name=sheet, engine="openpyxl", usecols=usecols
            )
            .reset_index(drop=False)
            .rename(columns={"index": "sheet_zero_based_row_index"})
        )
        df["sheet"] = sheet
        frames.append(df)
    combined = pd.concat(frames, ignore_index=True)
    mask = (
        combined["inchi_solute"].map(lambda x: supported_structure(x, "solute"))
        & combined["inchi_solvent1"].map(lambda x: supported_structure(x, "solvent"))
        & combined["T (K)"].notna()
    )
    if "inchi_solvent2" in combined.columns:
        mask &= combined["inchi_solvent2"].isna() | combined["inchi_solvent2"].map(
            lambda x: supported_structure(x, "solvent")
        )
    candidates = combined.loc[mask].copy().reset_index(drop=True)
    primary = []
    matched = []
    keep = []
    for _, row in candidates.iterrows():
        try:
            a = derive_functional_group_assignment(
                record_id="x",
                structure_identifier=str(row["inchi_solute"]),
                taxonomy=taxonomy,
                assignment_evidence="frozen taxonomy SMARTS on exact solute InChI before reading model outputs",
            )
            primary.append(a.primary_group)
            matched.append(list(a.matched_groups))
            keep.append(True)
        except Exception:  # noqa: BLE001 - taxonomy rejects structures broadly
            keep.append(False)
            primary.append(None)
            matched.append(None)
    candidates = candidates.loc[keep].copy().reset_index(drop=True)
    candidates["primary_functional_group"] = [x for x, k in zip(primary, keep) if k]
    candidates["all_matched_groups"] = [x for x, k in zip(matched, keep) if k]
    positions = np.unique(np.linspace(0, len(candidates) - 1, MAIN_SAMPLE_N, dtype=int))
    rows = [
        make_row_record(str(r["sheet"]), r, hm[str(r["sheet"])], int(pos))
        for pos, (_, r) in zip(
            positions, candidates.iloc[positions].iterrows(), strict=True
        )
    ]
    attach_labels_after_selection(rows)
    meta = {
        "selection_rule": "output-blind deterministic uniform sample from supported v1.1 active nonaqueous pure/binary sheets after adapter-domain and frozen-taxonomy filters; positions = unique(linspace(0,N-1,1000))",
        "candidate_rows_after_concat": len(combined),
        "candidate_rows_after_filters": len(candidates),
        "selected_row_count": len(rows),
        "distinct_primary_groups": sorted(
            {row["primary_functional_group"] for row in rows}
        ),
    }
    return rows, meta


def select_crosscheck_rows():
    hm = {sheet: header_map(sheet) for sheet in MAIN_SHEETS}
    pure = (
        pd.read_excel(
            WORKBOOK,
            sheet_name="From IDAC - no water",
            engine="openpyxl",
            usecols=[
                "inchi_solute",
                "DDB_name_solute",
                "inchi_solvent1",
                "DDB_name_solvent1",
                "T (K)",
                "Gsolv (kcal/mol)",
            ],
        )
        .reset_index(drop=False)
        .rename(columns={"index": "sheet_zero_based_row_index"})
    )
    pure = pure[
        pure["inchi_solute"].map(lambda x: supported_structure(x, "solute"))
        & pure["inchi_solvent1"].map(lambda x: supported_structure(x, "solvent"))
        & pure["T (K)"].notna()
    ]
    binary = (
        pd.read_excel(
            WORKBOOK,
            sheet_name="From VLE - Bin Solv - no water",
            engine="openpyxl",
            usecols=[
                "inchi_solute",
                "DDB_name_solute",
                "inchi_solvent1",
                "DDB_name_solvent1",
                "inchi_solvent2",
                "DDB_name_solvent2",
                "frac_solvent1",
                "T (K)",
                "Gsolv (kcal/mol)",
            ],
        )
        .reset_index(drop=False)
        .rename(columns={"index": "sheet_zero_based_row_index"})
    )
    binary = binary[
        binary["inchi_solute"].map(lambda x: supported_structure(x, "solute"))
        & binary["inchi_solvent1"].map(lambda x: supported_structure(x, "solvent"))
        & binary["inchi_solvent2"].map(lambda x: supported_structure(x, "solvent"))
        & binary["T (K)"].notna()
    ]
    rows = [
        make_row_record("From IDAC - no water", r, hm["From IDAC - no water"], int(i))
        for i, (_, r) in enumerate(pure.head(3).iterrows())
    ]
    rows += [
        make_row_record(
            "From VLE - Bin Solv - no water",
            r,
            hm["From VLE - Bin Solv - no water"],
            int(i),
        )
        for i, (_, r) in enumerate(binary.head(3).iterrows())
    ]
    return rows


def select_ternary_rows():
    hm = header_map(TERNARY_SHEET)
    df = (
        pd.read_excel(
            WORKBOOK,
            sheet_name=TERNARY_SHEET,
            engine="openpyxl",
            usecols=[
                "inchi_solute",
                "DDB_name_solute",
                "inchi_solvent1",
                "DDB_name_solvent1",
                "inchi_solvent2",
                "DDB_name_solvent2",
                "inchi_solvent3",
                "DDB_name_solvent3",
                "frac_solvent1",
                "frac_solvent2",
                "T (K)",
            ],
        )
        .reset_index(drop=False)
        .rename(columns={"index": "sheet_zero_based_row_index"})
    )
    mask = (
        df["inchi_solute"].map(lambda x: supported_structure(x, "solute"))
        & df["inchi_solvent1"].map(lambda x: supported_structure(x, "solvent"))
        & df["inchi_solvent2"].map(lambda x: supported_structure(x, "solvent"))
        & df["inchi_solvent3"].map(lambda x: supported_structure(x, "solvent"))
        & df["T (K)"].notna()
    )
    candidates = df.loc[mask].copy().reset_index(drop=True)
    positions = np.unique(
        np.linspace(
            0, len(candidates) - 1, min(TERNARY_SAMPLE_N, len(candidates)), dtype=int
        )
    )
    rows = [
        make_row_record(TERNARY_SHEET, r, hm, int(pos))
        for pos, (_, r) in zip(
            positions, candidates.iloc[positions].iterrows(), strict=True
        )
    ]
    attach_labels_after_selection(rows)
    meta = {
        "candidate_rows_after_filters": len(candidates),
        "selected_row_count": len(rows),
        "selection_rule": "output-blind deterministic uniform ternary diagnostic sample over supported rows",
    }
    return rows, meta


def score_rows(rows, cpu_inp, cpu_bundles, gpu_inp, gpu_bundles):
    cpu_preds = []
    t0 = time.perf_counter()
    for row in rows:
        cpu_preds.append(
            predict_direct(
                cpu_inp,
                cpu_bundles,
                solute_inchi=row["solute_inchi"],
                solvent_pairs=row["solvent_pairs"],
                temperature_k=row["temperature_kelvin"],
            )
        )
    cpu_seconds = time.perf_counter() - t0
    gpu_preds = []
    t1 = time.perf_counter()
    for row in rows:
        gpu_preds.append(
            predict_direct(
                gpu_inp,
                gpu_bundles,
                solute_inchi=row["solute_inchi"],
                solvent_pairs=row["solvent_pairs"],
                temperature_k=row["temperature_kelvin"],
            )
        )
    gpu_seconds = time.perf_counter() - t1
    cpu_errors = []
    gpu_errors = []
    max_pred_delta = 0.0
    max_model_delta = 0.0
    worsen_exact = []
    worsen_tol = []
    for i, row in enumerate(rows):
        cpu = cpu_preds[i]
        gpu = gpu_preds[i]
        exp = row["experimental_gsolv_kcal_mol"]
        cpu_signed = cpu["prediction_kcal_mol"] - exp
        gpu_signed = gpu["prediction_kcal_mol"] - exp
        cpu_abs = abs(cpu_signed)
        gpu_abs = abs(gpu_signed)
        pred_delta = gpu["prediction_kcal_mol"] - cpu["prediction_kcal_mol"]
        max_pred_delta = max(max_pred_delta, abs(pred_delta))
        model_deltas = []
        for c, g in zip(
            cpu["model_predictions"], gpu["model_predictions"], strict=True
        ):
            delta = g["g_temperature_kcal_mol"] - c["g_temperature_kcal_mol"]
            max_model_delta = max(max_model_delta, abs(delta))
            model_deltas.append(
                {
                    "model_index": c["model_index"],
                    "cpu_g_temperature_kcal_mol": c["g_temperature_kcal_mol"],
                    "gpu_g_temperature_kcal_mol": g["g_temperature_kcal_mol"],
                    "gpu_minus_cpu_kcal_mol": delta,
                    "abs_gpu_minus_cpu_kcal_mol": abs(delta),
                }
            )
        if gpu_abs > cpu_abs:
            worsen_exact.append(row["row_id"])
        if gpu_abs > cpu_abs + GPU_WORSEN_TOL:
            worsen_tol.append(row["row_id"])
        row.update(
            {
                "canonical_solute_smiles": cpu["canonical_solute_smiles"],
                "canonical_solvent_components": cpu["canonical_solvents"],
                "cpu_prediction_kcal_mol": cpu["prediction_kcal_mol"],
                "cpu_signed_error_kcal_mol": cpu_signed,
                "cpu_abs_error_kcal_mol": cpu_abs,
                "gpu_prediction_kcal_mol": gpu["prediction_kcal_mol"],
                "gpu_signed_error_kcal_mol": gpu_signed,
                "gpu_abs_error_kcal_mol": gpu_abs,
                "gpu_minus_cpu_prediction_kcal_mol": pred_delta,
                "abs_gpu_minus_cpu_prediction_kcal_mol": abs(pred_delta),
                "gpu_worsens_abs_error_exact": gpu_abs > cpu_abs,
                "gpu_worsens_abs_error_over_5e12": gpu_abs > cpu_abs + GPU_WORSEN_TOL,
                "max_abs_gpu_minus_cpu_model_kcal_mol": max(
                    x["abs_gpu_minus_cpu_kcal_mol"] for x in model_deltas
                ),
                "model_deltas": model_deltas,
                "live_cpu_minus_workbook_qmexp_kcal_mol": (
                    cpu["prediction_kcal_mol"]
                    - row["published_workbook_qmexp_gt_kcal_mol"]
                    if row["published_workbook_qmexp_gt_kcal_mol"] is not None
                    else None
                ),
            }
        )
        cpu_errors.append(cpu_signed)
        gpu_errors.append(gpu_signed)
    return {
        "cpu_seconds": cpu_seconds,
        "gpu_seconds": gpu_seconds,
        "cpu_metrics": metrics(cpu_errors),
        "gpu_metrics": metrics(gpu_errors),
        "gpu_vs_cpu": {
            "gpu_worsens_abs_error_exact_count": len(worsen_exact),
            "gpu_worsens_abs_error_over_5e12_count": len(worsen_tol),
            "gpu_worsen_exact_row_ids": worsen_exact,
            "gpu_worsen_over_5e12_row_ids": worsen_tol,
            "max_abs_gpu_minus_cpu_prediction_kcal_mol": max_pred_delta,
            "max_abs_gpu_minus_cpu_model_kcal_mol": max_model_delta,
        },
    }


def crosscheck(cpu_inp, cpu_bundles, gpu_inp, gpu_bundles):
    adapter = SolPropMixPropertyAdapter(
        SOURCE_ROOT, STATIC_ROOT, WEIGHTS_ROOT, python_executable=PYTHON_EXE
    )
    rows = select_crosscheck_rows()
    checks = []
    for row in rows:
        for device_name, inp, bundles in (
            ("cpu", cpu_inp, cpu_bundles),
            ("cuda:0", gpu_inp, gpu_bundles),
        ):
            direct = predict_direct(
                inp,
                bundles,
                solute_inchi=row["solute_inchi"],
                solvent_pairs=row["solvent_pairs"],
                temperature_k=row["temperature_kelvin"],
            )
            result = adapter.predict(
                row["solute_inchi"],
                row["solvent_pairs"],
                temperature_kelvin=row["temperature_kelvin"],
                device=device_name,
                precision="float64_promoted",
            )
            pred_delta = abs(
                result.predicted_solvation_free_energy_kcal_mol
                - direct["prediction_kcal_mol"]
            )
            model_deltas = [
                {
                    "model_index": model_index,
                    "prediction_abs_delta_kcal_mol": abs(
                        adapter_model.g_temperature_kcal_mol
                        - direct_model["g_temperature_kcal_mol"]
                    ),
                }
                for model_index, (adapter_model, direct_model) in enumerate(
                    zip(
                        result.model_predictions,
                        direct["model_predictions"],
                        strict=True,
                    )
                )
            ]
            model_delta = max(
                item["prediction_abs_delta_kcal_mol"] for item in model_deltas
            )
            checks.append(
                {
                    "row_id": row["row_id"],
                    "device": device_name,
                    "prediction_abs_delta_kcal_mol": pred_delta,
                    "model_prediction_abs_deltas_kcal_mol": model_deltas,
                    "max_model_prediction_abs_delta_kcal_mol": model_delta,
                    "passes_5e12": pred_delta <= GPU_WORSEN_TOL
                    and model_delta <= GPU_WORSEN_TOL,
                }
            )
    return checks


def _csv_row(row: dict[str, Any]) -> dict[str, Any]:
    return {
        key: (
            json.dumps(value, ensure_ascii=False)
            if isinstance(value, (list, dict))
            else value
        )
        for key, value in row.items()
        if key != "model_deltas"
    }


def _panel_record(row: dict[str, Any]) -> dict[str, Any]:
    record = {
        key: row[key]
        for key in (
            "row_id",
            "sheet",
            "excel_row_number",
            "sheet_zero_based_row_index",
            "selection_position",
            "primary_functional_group",
            "all_matched_groups",
            "solute_name",
            "solute_inchi",
            "temperature_kelvin",
            "solvent_pairs",
            "solvent_names",
            "input_cells",
            "canonical_solute_smiles",
            "canonical_solvent_components",
        )
    }
    record["input_cells"] = {
        key: value
        for key, value in record["input_cells"].items()
        if key != "experimental_gsolv"
    }
    return record


def _reference_record(row: dict[str, Any]) -> dict[str, Any]:
    cells = row["input_cells"]
    return {
        "row_id": row["row_id"],
        "experimental_gsolv_kcal_mol": row["experimental_gsolv_kcal_mol"],
        "experimental_value_cell": cells["experimental_gsolv"],
        "published_workbook_qmexp_h298_kcal_mol": row[
            "published_workbook_qmexp_h298_kcal_mol"
        ],
        "published_workbook_qmexp_g298_kcal_mol": row[
            "published_workbook_qmexp_g298_kcal_mol"
        ],
        "published_workbook_qmexp_gt_kcal_mol": row[
            "published_workbook_qmexp_gt_kcal_mol"
        ],
    }


def _add_canonical_inputs(rows: list[dict[str, Any]]) -> None:
    for row in rows:
        solute, solvents, _ = prepare_request(row["solute_inchi"], row["solvent_pairs"])
        row["canonical_solute_smiles"] = solute
        row["canonical_solvent_components"] = [
            {
                "canonical_smiles": item.canonical_smiles,
                "mole_fraction": item.mole_fraction,
            }
            for item in solvents
        ]


def _read_csv_records(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def _write_csv_records(path: Path, rows: list[dict[str, str]]) -> None:
    if not rows:
        raise RuntimeError(f"cannot write empty frozen CSV: {path.name}")
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]), lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def _frozen_score(rows: list[dict[str, str]]) -> dict[str, Any]:
    cpu_errors = [float(row["cpu_signed_error_kcal_mol"]) for row in rows]
    gpu_errors = [float(row["gpu_signed_error_kcal_mol"]) for row in rows]
    exact_ids = [
        row["row_id"]
        for row in rows
        if abs(float(row["gpu_signed_error_kcal_mol"]))
        > abs(float(row["cpu_signed_error_kcal_mol"]))
    ]
    tolerance_ids = [
        row["row_id"]
        for row in rows
        if abs(float(row["gpu_signed_error_kcal_mol"]))
        > abs(float(row["cpu_signed_error_kcal_mol"])) + GPU_WORSEN_TOL
    ]
    return {
        "cpu_metrics": metrics(cpu_errors),
        "gpu_metrics": metrics(gpu_errors),
        "gpu_vs_cpu": {
            "gpu_worsens_abs_error_exact_count": len(exact_ids),
            "gpu_worsens_abs_error_over_5e12_count": len(tolerance_ids),
            "gpu_worsen_exact_row_ids": exact_ids,
            "gpu_worsen_over_5e12_row_ids": tolerance_ids,
            "max_abs_gpu_minus_cpu_prediction_kcal_mol": max(
                abs(float(row["gpu_minus_cpu_prediction_kcal_mol"])) for row in rows
            ),
            "max_abs_gpu_minus_cpu_model_kcal_mol": max(
                float(row["max_abs_gpu_minus_cpu_model_kcal_mol"]) for row in rows
            ),
        },
    }


def _typed_result_rows(rows: list[dict[str, str]]) -> list[dict[str, Any]]:
    return [
        {
            "row_id": row["row_id"],
            "primary_functional_group": row["primary_functional_group"],
            "cpu_prediction_kcal_mol": float(row["cpu_prediction_kcal_mol"]),
            "gpu_prediction_kcal_mol": float(row["gpu_prediction_kcal_mol"]),
            "cpu_signed_error_kcal_mol": float(row["cpu_signed_error_kcal_mol"]),
            "gpu_signed_error_kcal_mol": float(row["gpu_signed_error_kcal_mol"]),
            "live_cpu_minus_workbook_qmexp_kcal_mol": float(
                row["live_cpu_minus_workbook_qmexp_kcal_mol"]
            ),
        }
        for row in rows
    ]


def _assert_float_field(
    row: dict[str, str], field: str, expected: float, *, row_id: str
) -> None:
    if float(row[field]) != expected:
        raise RuntimeError(f"{row_id}: inconsistent {field}")


def _validate_scored_rows(
    raw_rows: list[dict[str, str]], selected_rows: list[dict[str, Any]]
) -> None:
    if len(raw_rows) != len(selected_rows):
        raise RuntimeError("frozen scored row count does not match selection")
    for raw, selected in zip(raw_rows, selected_rows, strict=True):
        row_id = selected["row_id"]
        for field in (
            "row_id",
            "sheet",
            "solute_name",
            "solute_inchi",
            "canonical_solute_smiles",
        ):
            if raw[field] != str(selected[field]):
                raise RuntimeError(f"{row_id}: inconsistent {field}")
        expected_group = selected["primary_functional_group"]
        if raw["primary_functional_group"] != (
            "" if expected_group is None else str(expected_group)
        ):
            raise RuntimeError(f"{row_id}: inconsistent primary_functional_group")
        for field in (
            "excel_row_number",
            "sheet_zero_based_row_index",
            "selection_position",
        ):
            if int(raw[field]) != selected[field]:
                raise RuntimeError(f"{row_id}: inconsistent {field}")
        for field in (
            "temperature_kelvin",
            "experimental_gsolv_kcal_mol",
            "published_workbook_qmexp_h298_kcal_mol",
            "published_workbook_qmexp_g298_kcal_mol",
            "published_workbook_qmexp_gt_kcal_mol",
        ):
            _assert_float_field(raw, field, float(selected[field]), row_id=row_id)
        expected_groups = selected["all_matched_groups"]
        if raw["all_matched_groups"] != (
            "" if expected_groups is None else json.dumps(expected_groups)
        ):
            raise RuntimeError(f"{row_id}: inconsistent all_matched_groups")
        for field in (
            "solvent_pairs",
            "solvent_names",
            "input_cells",
            "canonical_solvent_components",
        ):
            expected = json.loads(json.dumps(selected[field]))
            if json.loads(raw[field]) != expected:
                raise RuntimeError(f"{row_id}: inconsistent {field}")

        experiment = float(raw["experimental_gsolv_kcal_mol"])
        cpu_prediction = float(raw["cpu_prediction_kcal_mol"])
        gpu_prediction = float(raw["gpu_prediction_kcal_mol"])
        cpu_error = cpu_prediction - experiment
        gpu_error = gpu_prediction - experiment
        prediction_delta = gpu_prediction - cpu_prediction
        _assert_float_field(raw, "cpu_signed_error_kcal_mol", cpu_error, row_id=row_id)
        _assert_float_field(raw, "gpu_signed_error_kcal_mol", gpu_error, row_id=row_id)
        _assert_float_field(
            raw, "cpu_abs_error_kcal_mol", abs(cpu_error), row_id=row_id
        )
        _assert_float_field(
            raw, "gpu_abs_error_kcal_mol", abs(gpu_error), row_id=row_id
        )
        _assert_float_field(
            raw,
            "gpu_minus_cpu_prediction_kcal_mol",
            prediction_delta,
            row_id=row_id,
        )
        _assert_float_field(
            raw,
            "abs_gpu_minus_cpu_prediction_kcal_mol",
            abs(prediction_delta),
            row_id=row_id,
        )
        exact_worsening = abs(gpu_error) > abs(cpu_error)
        tolerance_worsening = abs(gpu_error) > abs(cpu_error) + GPU_WORSEN_TOL
        if (raw["gpu_worsens_abs_error_exact"] == "True") != exact_worsening:
            raise RuntimeError(f"{row_id}: inconsistent exact GPU worsening flag")
        if (raw["gpu_worsens_abs_error_over_5e12"] == "True") != tolerance_worsening:
            raise RuntimeError(f"{row_id}: inconsistent tolerance GPU worsening flag")
        live_delta = cpu_prediction - float(raw["published_workbook_qmexp_gt_kcal_mol"])
        _assert_float_field(
            raw,
            "live_cpu_minus_workbook_qmexp_kcal_mol",
            live_delta,
            row_id=row_id,
        )


def _validate_crosscheck(checks: list[dict[str, Any]]) -> None:
    row_ids = [
        "From IDAC - no water!2",
        "From IDAC - no water!3",
        "From IDAC - no water!4",
        "From VLE - Bin Solv - no water!2",
        "From VLE - Bin Solv - no water!3",
        "From VLE - Bin Solv - no water!4",
    ]
    expected_pairs = [
        (row_id, device) for row_id in row_ids for device in ("cpu", "cuda:0")
    ]
    if [(item.get("row_id"), item.get("device")) for item in checks] != expected_pairs:
        raise RuntimeError("frozen adapter crosscheck row/device design mismatch")
    expected_fields = {
        "row_id",
        "device",
        "prediction_abs_delta_kcal_mol",
        "model_prediction_abs_deltas_kcal_mol",
        "max_model_prediction_abs_delta_kcal_mol",
        "passes_5e12",
    }
    for item in checks:
        if set(item) != expected_fields:
            raise RuntimeError("frozen adapter crosscheck field set mismatch")
        model_deltas = item["model_prediction_abs_deltas_kcal_mol"]
        if [entry.get("model_index") for entry in model_deltas] != list(range(10)):
            raise RuntimeError("frozen adapter crosscheck model index mismatch")
        if any(
            set(entry) != {"model_index", "prediction_abs_delta_kcal_mol"}
            for entry in model_deltas
        ):
            raise RuntimeError("frozen adapter model-delta field set mismatch")
        deltas = [
            float(entry["prediction_abs_delta_kcal_mol"]) for entry in model_deltas
        ]
        prediction_delta = float(item["prediction_abs_delta_kcal_mol"])
        maximum = max(deltas)
        if prediction_delta != 0.0 or maximum != 0.0:
            raise RuntimeError(
                "frozen adapter crosscheck must retain exact zero deltas"
            )
        if float(item["max_model_prediction_abs_delta_kcal_mol"]) != maximum:
            raise RuntimeError("frozen adapter crosscheck maximum is inconsistent")
        expected_pass = prediction_delta <= GPU_WORSEN_TOL and maximum <= GPU_WORSEN_TOL
        if item["passes_5e12"] is not expected_pass:
            raise RuntimeError("frozen adapter crosscheck pass flag is inconsistent")


def assemble_from_frozen(frozen_dir: Path) -> None:
    """Rebuild canonical artifacts without importing or loading upstream runtime."""
    panel = json.loads((frozen_dir / "panel-inputs.json").read_text())
    references = json.loads((frozen_dir / "experimental-references.json").read_text())
    raw_main = _read_csv_records(frozen_dir / "rows.csv")
    raw_ternary = _read_csv_records(frozen_dir / "ternary-rows.csv")
    checks = json.loads((frozen_dir / "adapter-crosscheck.json").read_text())

    selected_main, main_meta = select_main_rows()
    selected_ternary, ternary_meta = select_ternary_rows()
    _add_canonical_inputs(selected_main)
    _add_canonical_inputs(selected_ternary)
    expected_panel = json.loads(
        json.dumps(
            {
                "model_worker_may_read_only_this_file": True,
                "selection_labels_read_after_indices_frozen": True,
                "records": [_panel_record(row) for row in selected_main],
            }
        )
    )
    expected_references = {
        "must_not_be_supplied_to_model_worker": True,
        "records": [_reference_record(row) for row in selected_main],
    }
    if panel != expected_panel:
        for index, (actual, expected) in enumerate(
            zip(panel["records"], expected_panel["records"], strict=True)
        ):
            if actual != expected:
                differing = sorted(
                    key
                    for key in set(actual) | set(expected)
                    if actual.get(key) != expected.get(key)
                )
                raise RuntimeError(
                    "frozen panel does not match official v1.1 selection: "
                    f"record {index}, fields {differing}"
                )
        raise RuntimeError(
            "frozen panel envelope does not match official v1.1 selection"
        )
    if references != expected_references:
        raise RuntimeError("frozen references do not match official v1.1 labels")
    main_ids = [row["row_id"] for row in raw_main]
    ternary_ids = [row["row_id"] for row in raw_ternary]
    if main_ids != [row["row_id"] for row in selected_main]:
        raise RuntimeError("frozen main CSV is not bound to selected panel")
    if ternary_ids != [row["row_id"] for row in selected_ternary]:
        raise RuntimeError("frozen ternary CSV is not bound to selected diagnostic")
    _validate_scored_rows(raw_main, selected_main)
    _validate_scored_rows(raw_ternary, selected_ternary)
    _validate_crosscheck(checks)

    main_rows = _typed_result_rows(raw_main)
    ternary_rows = _typed_result_rows(raw_ternary)
    main_score = _frozen_score(raw_main)
    ternary_score = _frozen_score(raw_ternary)
    by_group: dict[str, Any] = {}
    for group in sorted({row["primary_functional_group"] for row in raw_main}):
        subset = [row for row in raw_main if row["primary_functional_group"] == group]
        by_group[group] = {
            "n": len(subset),
            "cpu": metrics([float(row["cpu_signed_error_kcal_mol"]) for row in subset]),
            "gpu": metrics([float(row["gpu_signed_error_kcal_mol"]) for row in subset]),
        }

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    panel_path = OUT_DIR / "panel-inputs.json"
    reference_path = OUT_DIR / "experimental-references.json"
    rows_path = OUT_DIR / "rows.csv"
    ternary_path = OUT_DIR / "ternary-rows.csv"
    crosscheck_path = OUT_DIR / "adapter-crosscheck.json"
    panel_path.write_text(json.dumps(panel, indent=2, ensure_ascii=False) + "\n")
    reference_path.write_text(
        json.dumps(references, indent=2, ensure_ascii=False) + "\n"
    )
    _write_csv_records(rows_path, raw_main)
    _write_csv_records(ternary_path, raw_ternary)
    crosscheck_path.write_text(json.dumps(checks, indent=2) + "\n")
    paths = (panel_path, reference_path, rows_path, ternary_path, crosscheck_path)
    result = _canonical_result(
        main_rows,
        ternary_rows,
        main_meta,
        ternary_meta,
        main_score,
        ternary_score,
        by_group,
        checks,
        {path.name: sha256_file(path) for path in paths},
        construction_mode="exact_canonical_reconstruction_from_fixed_scored_inputs",
    )
    result_path = OUT_DIR / "result.json"
    result_path.write_text(json.dumps(result, indent=2, ensure_ascii=False) + "\n")


def _canonical_result(
    main_rows: list[dict[str, Any]],
    ternary_rows: list[dict[str, Any]],
    main_meta: dict[str, Any],
    ternary_meta: dict[str, Any],
    main_score: dict[str, Any],
    ternary_score: dict[str, Any],
    by_group: dict[str, Any],
    checks: list[dict[str, Any]],
    artifact_hashes: dict[str, str],
    *,
    construction_mode: str,
) -> dict[str, Any]:
    main_mismatches = sum(
        row["cpu_prediction_kcal_mol"] != row["gpu_prediction_kcal_mol"]
        for row in main_rows
    )
    ternary_mismatches = sum(
        row["cpu_prediction_kcal_mol"] != row["gpu_prediction_kcal_mol"]
        for row in ternary_rows
    )
    return {
        "schema_version": 1,
        "artifact_construction": {
            "mode": construction_mode,
            "not_a_new_live_run": construction_mode
            == "exact_canonical_reconstruction_from_fixed_scored_inputs",
            "fixed_frozen_input_sha256": (
                CANONICAL_FROZEN_INPUT_SHA256
                if construction_mode
                == "exact_canonical_reconstruction_from_fixed_scored_inputs"
                else None
            ),
        },
        "scientific_identity": {
            "source_repository": SOLPROPMIX_SOURCE_URL,
            "source_revision": SOLPROPMIX_SOURCE_REVISION,
            "source_tree": SOLPROPMIX_SOURCE_TREE,
            "checkpoint_family": SOLPROPMIX_CHECKPOINT_FAMILY,
            "data_release_version": SOLPROPMIX_DATA_RELEASE_VERSION,
            "data_release_record": SOLPROPMIX_DATA_RELEASE_RECORD,
            "previous_data_release_record": SOLPROPMIX_PREVIOUS_DATA_RELEASE_RECORD,
            "supplemental_source_origin": SOLPROPMIX_SUPPLEMENTAL_SOURCE_ORIGIN,
            "checkpoint_origin": SOLPROPMIX_CHECKPOINT_ORIGIN,
            "checkpoints_byte_identical_across_v1_0_v1_1": SOLPROPMIX_CHECKPOINTS_BYTE_IDENTICAL_ACROSS_V1_0_V1_1,
            "workbook_member": "Data/solprop-mix_v1.1.xlsx",
            "workbook_sha256": WORKBOOK_SHA256,
            "static_code_zip": {
                "name": SOLPROPMIX_STATIC_CODE_ZIP_NAME,
                "size_bytes": SOLPROPMIX_STATIC_CODE_ZIP_SIZE_BYTES,
                "md5": SOLPROPMIX_STATIC_CODE_ZIP_MD5,
                "sha256": SOLPROPMIX_STATIC_CODE_ZIP_SHA256,
            },
            "modelweights_zip": {
                "size_bytes": SOLPROPMIX_MODELWEIGHTS_ZIP_SIZE_BYTES,
                "md5": SOLPROPMIX_MODELWEIGHTS_ZIP_MD5,
                "sha256": SOLPROPMIX_MODELWEIGHTS_ZIP_SHA256,
            },
            "checkpoint_sha256": list(SOLPROPMIX_CHECKPOINT_SHA256),
            "static_module_sha256": OFFICIAL_STATIC_PYTHON_SHA256,
            "taxonomy_sha256": TAXONOMY_SHA256,
            "runner_relative_path": "docs/pretrained-solvation-hub/run_solpropmix_qmexp_broad_audit.py",
            "runner_sha256": sha256_file(Path(__file__)),
            "adapter_relative_path": "maple/function/solvfe/solpropmix_property.py",
            "adapter_sha256": sha256_file(
                REPO / "maple/function/solvfe/solpropmix_property.py"
            ),
            "worker_relative_path": "maple/function/solvfe/_solpropmix_worker.py",
            "worker_sha256": sha256_file(
                REPO / "maple/function/solvfe/_solpropmix_worker.py"
            ),
        },
        "selection_protocol": {
            "label_separation": "read input columns, freeze deterministic row indices, then separately read labels by frozen sheet row",
            "model_output_blind": True,
            "deterministic": True,
            "main": main_meta,
            "ternary": ternary_meta,
        },
        "strict_float64_policy": {
            "execution_dtype": "float64",
            "deterministic_algorithms": True,
            "tf32": False,
            "amp": False,
            "fp16": False,
            "bf16": False,
            "diagnostic_tolerance_kcal_mol": GPU_WORSEN_TOL,
            "zero_loss_requires_exact_no_error_worsening": True,
        },
        "main_supported_lane": {
            "sheets": list(MAIN_SHEETS),
            "overall_cpu": main_score["cpu_metrics"],
            "overall_gpu": main_score["gpu_metrics"],
            "by_primary_functional_group": by_group,
            "gpu_vs_cpu": {
                **main_score["gpu_vs_cpu"],
                "ensemble_prediction_bitwise_mismatch_count": main_mismatches,
                "strict_zero_loss": False,
                "gpu_admission": False,
                "reason": "exact absolute error worsened on at least one row despite no delta exceeding the 5e-12 diagnostic tolerance",
            },
            "accuracy_gate": {
                "mae_threshold_kcal_mol": 0.25,
                "rmse_threshold_kcal_mol": 0.37,
                "mae_pass": main_score["cpu_metrics"]["mae_kcal_mol"] <= 0.25,
                "rmse_pass": main_score["cpu_metrics"]["rmse_kcal_mol"] <= 0.37,
                "overall_pass": False,
            },
            "training_holdout_boundary": "training overlap is unknown in public assets; this is a development/non-blind benchmark, not a final blind holdout",
            "matched_qm_speed_eligible": False,
            "speed_boundary": "accuracy gates fail, so wall-clock timing cannot be used for admission",
        },
        "adapter_crosscheck": {
            "rows": checks,
            "row_design": "three pure plus three binary rows, each on CPU and GPU",
            "threshold_kcal_mol": GPU_WORSEN_TOL,
            "all_per_model_and_ensemble_pass": all(
                item["passes_5e12"] for item in checks
            ),
        },
        "ternary_undocumented_dynamic_slot_diagnostic": {
            "sheet": TERNARY_SHEET,
            "not_part_of_main_lane_or_admission": True,
            "selection": ternary_meta,
            "overall_cpu": ternary_score["cpu_metrics"],
            "overall_gpu": ternary_score["gpu_metrics"],
            "gpu_vs_cpu": {
                **ternary_score["gpu_vs_cpu"],
                "ensemble_prediction_bitwise_mismatch_count": ternary_mismatches,
            },
            "mean_abs_live_cpu_vs_workbook_qmexp_gt_kcal_mol": math.fsum(
                abs(row["live_cpu_minus_workbook_qmexp_kcal_mol"])
                for row in ternary_rows
            )
            / len(ternary_rows),
            "max_abs_live_cpu_vs_workbook_qmexp_gt_kcal_mol": max(
                abs(row["live_cpu_minus_workbook_qmexp_kcal_mol"])
                for row in ternary_rows
            ),
        },
        "artifact_sha256": artifact_hashes,
        "canonical_result_exclusions": [
            "wall-clock timings",
            "absolute cache paths",
            "stdout or stderr stream hashes",
            "redundant per-model row JSON",
        ],
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", required=True, type=Path)
    parser.add_argument("--source-root", required=True, type=Path)
    parser.add_argument("--static-root", required=True, type=Path)
    parser.add_argument("--weights-root", required=True, type=Path)
    parser.add_argument("--python-executable", required=True, type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    parser.add_argument(
        "--assemble-from-frozen",
        type=Path,
        help="rebuild artifacts from scored frozen CSVs after official selection checks",
    )
    return parser.parse_args()


def main() -> None:
    global WORKBOOK, SOURCE_ROOT, STATIC_ROOT, WEIGHTS_ROOT, PYTHON_EXE, OUT_DIR
    args = parse_args()
    WORKBOOK = args.workbook.resolve()
    SOURCE_ROOT = args.source_root.resolve()
    STATIC_ROOT = args.static_root.resolve()
    WEIGHTS_ROOT = args.weights_root.resolve()
    PYTHON_EXE = args.python_executable.resolve()
    OUT_DIR = args.output_dir.resolve()
    frozen_snapshot = (
        snapshot_canonical_frozen_inputs(args.assemble_from_frozen.resolve())
        if args.assemble_from_frozen is not None
        else None
    )
    authenticate_official_assets(
        WORKBOOK, SOURCE_ROOT, STATIC_ROOT, WEIGHTS_ROOT, PYTHON_EXE
    )
    WORKBOOK, WEIGHTS_ROOT = snapshot_official_data(WORKBOOK, WEIGHTS_ROOT)
    if frozen_snapshot is not None:
        assemble_from_frozen(frozen_snapshot)
        result_path = OUT_DIR / "result.json"
        print(
            json.dumps(
                {
                    "result": str(result_path),
                    "result_sha256": sha256_file(result_path),
                },
                indent=2,
            )
        )
        return

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    configure_runtime(STATIC_ROOT)
    main_rows, main_meta = select_main_rows()
    ternary_rows, ternary_meta = select_ternary_rows()
    cpu_inp, cpu_bundles, _ = build_models("cpu")
    gpu_inp, gpu_bundles, _ = build_models("cuda:0")
    checks = crosscheck(cpu_inp, cpu_bundles, gpu_inp, gpu_bundles)
    if not all(item["passes_5e12"] for item in checks):
        raise RuntimeError("adapter crosscheck failed")
    main_score = score_rows(main_rows, cpu_inp, cpu_bundles, gpu_inp, gpu_bundles)
    ternary_score = score_rows(ternary_rows, cpu_inp, cpu_bundles, gpu_inp, gpu_bundles)
    by_group = {}
    for group in sorted({row["primary_functional_group"] for row in main_rows}):
        subset = [row for row in main_rows if row["primary_functional_group"] == group]
        by_group[group] = {
            "n": len(subset),
            "cpu": metrics([row["cpu_signed_error_kcal_mol"] for row in subset]),
            "gpu": metrics([row["gpu_signed_error_kcal_mol"] for row in subset]),
        }

    panel_path = OUT_DIR / "panel-inputs.json"
    reference_path = OUT_DIR / "experimental-references.json"
    rows_path = OUT_DIR / "rows.csv"
    ternary_path = OUT_DIR / "ternary-rows.csv"
    crosscheck_path = OUT_DIR / "adapter-crosscheck.json"
    result_path = OUT_DIR / "result.json"
    panel_path.write_text(
        json.dumps(
            {
                "model_worker_may_read_only_this_file": True,
                "selection_labels_read_after_indices_frozen": True,
                "records": [_panel_record(row) for row in main_rows],
            },
            indent=2,
            ensure_ascii=False,
        )
        + "\n"
    )
    reference_path.write_text(
        json.dumps(
            {
                "must_not_be_supplied_to_model_worker": True,
                "records": [_reference_record(row) for row in main_rows],
            },
            indent=2,
            ensure_ascii=False,
        )
        + "\n"
    )
    pd.DataFrame([_csv_row(row) for row in main_rows]).to_csv(rows_path, index=False)
    pd.DataFrame([_csv_row(row) for row in ternary_rows]).to_csv(
        ternary_path, index=False
    )
    crosscheck_path.write_text(json.dumps(checks, indent=2) + "\n")
    artifact_hashes = {
        path.name: sha256_file(path)
        for path in (
            panel_path,
            reference_path,
            rows_path,
            ternary_path,
            crosscheck_path,
        )
    }
    result = _canonical_result(
        main_rows,
        ternary_rows,
        main_meta,
        ternary_meta,
        main_score,
        ternary_score,
        by_group,
        checks,
        artifact_hashes,
        construction_mode="live_official_asset_execution",
    )
    result_path.write_text(json.dumps(result, indent=2, ensure_ascii=False) + "\n")
    print(
        json.dumps(
            {"result": str(result_path), "result_sha256": sha256_file(result_path)},
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
